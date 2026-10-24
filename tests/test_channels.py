"""Channels: a letter's attachments become documents, the owner's tables export, the training fines portal and the
robot's reconciliation, VK and MAX on the bot's core, and documents with personal data kept off the cloud."""

import io
import re
from email.message import EmailMessage

import pytest
from sqlalchemy import select

from app import channels, gov, inbox, robot
from app.db import Item, Job, db
from app.seed import add
from tests.test_integration import client  # noqa: F401  (fixture)


def letter_with_pdf() -> bytes:
    msg = EmailMessage()
    msg["From"] = "СТО <sto@example.org>"
    msg["To"] = "docs@fleet.example"
    msg["Subject"] = "Счёт"
    msg.set_content("Во вложении счёт.")
    msg.add_attachment(b"%PDF-1.4\n%training\n", maintype="application", subtype="pdf", filename="schet.pdf")
    msg.add_attachment(b"not a document", maintype="application", subtype="octet-stream", filename="x.bin")
    return bytes(msg)


def test_letter_attachments_become_documents(client):  # noqa: F811
    c, sid = client
    r = c.post("/api/v1/inbox/eml", files={"file": ("letter.eml", letter_with_pdf(), "message/rfc822")})
    assert r.status_code == 200, r.text
    docs = r.json()["docs"]
    assert len(docs) == 1 and docs[0]["source"] == "mail" and docs[0]["mail"]["from"] == "sto@example.org"
    with db() as s:
        assert s.scalar(select(Job).where(Job.space == sid, Job.kind == "inbox").order_by(Job.due.desc()))
    empty = EmailMessage()
    empty["From"] = "x@example.org"
    empty.set_content("коротко")
    assert c.post("/api/v1/inbox/eml", files={"file": ("e.eml", bytes(empty), "message/rfc822")}).status_code == 400


def test_xlsx_export_has_three_sheets(client):  # noqa: F811
    from openpyxl import load_workbook

    c, _ = client
    r = c.get("/api/v1/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Сводка", "Сделать сейчас", "Документы"]
    assert wb["Документы"]["A1"].value == "Код"
    assert c.post("/api/v1/export/sheets").status_code == 409  # no service account: says so, XLSX still works


def test_training_portal_behaves_like_a_portal(client):  # noqa: F811
    c, _ = client
    page = c.get("/gov").text
    cid = re.search(r'name="captcha_id" value="(\w+)"', page).group(1)
    code = re.search(r'data-test-code="(\d+)"', page).group(1)
    wrong = c.post("/gov/search", data={"plates": "ДЕМО 0002", "captcha": "0", "captcha_id": cid}, follow_redirects=False)
    assert wrong.headers["location"] == "/gov?error=captcha"
    page = c.get("/gov").text  # a code works once: a new one
    cid = re.search(r'name="captcha_id" value="(\w+)"', page).group(1)
    code = re.search(r'data-test-code="(\d+)"', page).group(1)
    plates = [f"ДЕМО {i:04d}" for i in range(1, 25)] + ["<script>alert(1)</script> 0002"]
    found = c.post("/gov/search", data={"plates": "\n".join(plates), "captcha": code, "captcha_id": cid}, follow_redirects=False)
    url = found.headers["location"]
    first = c.get(url)
    assert first.status_code == 200 and "DEMO-FINE-TP-0002" in first.text
    assert c.get(url).status_code == 429  # faster than 0.3 s
    import time

    time.sleep(0.35)
    second = url.replace("page=1", "page=2")
    assert c.get(second).status_code == 503  # the first ask of page 2: the portal is under load
    time.sleep(0.35)
    assert c.get(second).status_code == 200
    time.sleep(0.35)
    last = c.get(url.replace("page=1", "page=3")).text
    assert "<script>" not in last and "&lt;SCRIPT&gt;" in last


def test_robot_reconciles_and_loads(client):  # noqa: F811
    c, sid = client
    rows = [f for p in ("ДЕМО 0002", "ДЕМО 0011") for f in gov.fines_for(p)]
    with db() as s:
        # what n8n delivered: TP-0002 with the same amount, TP-0011 with another, TP-0026 the portal no longer shows
        for code, amount in (("TP-0002", "500.00"), ("TP-0011", "500.00"), ("TP-0026", "500.00")):
            add(s, sid, "source_record", "fines:" + code, {
                "source": "fines", "values": {"vehicle_code": code, "fine": amount, "external_id": "DEMO-FINE-" + code}})
        s.flush()
        rec = robot.reconcile(s, sid, rows)
    states = {r["resolution"]: r["state"] for r in rec}
    assert states["DEMO-FINE-TP-0002"] == "ok" and states["DEMO-FINE-TP-0011"] == "amount"
    assert states["DEMO-FINE-TP-0026"] == "ledger-only"
    # A run whose portal-only fines are loaded into the ledger as the records n8n writes.
    extra = {"resolution": "18810177000000000001", "plate": "ДЕМО 0001", "amount": 1500, "date": "2026-10-20",
             "discount_until": "2026-11-09", "state": "portal-only", "note": ""}
    with db() as s:
        run = add(s, sid, "robot_run", "ROBOT-T", {"status": "done", "reconciliation": [extra]})
        s.flush()
        rid = run.id
    loaded = c.post(f"/api/v1/robot/{rid}/load").json()
    assert loaded["loaded"] == 1 and loaded["run"]["reconciliation"][0]["state"] == "loaded"
    assert c.post(f"/api/v1/robot/{rid}/load").json()["loaded"] == 0  # once
    with db() as s:
        rec = s.scalar(select(Item).where(Item.space == sid, Item.code == "fines:robot:18810177000000000001"))
        assert rec.data["values"]["via"] == "robot" and rec.data["values"]["fine"] == "1500.00"


def test_vk_and_max_on_the_bot_core(client, monkeypatch):  # noqa: F811
    c, _ = client
    vk = channels.from_vk({"type": "message_event", "event_id": "e1",
                           "object": {"user_id": 7, "peer_id": 7, "event_id": "x", "payload": {"cmd": "role:driver"}}})
    assert vk["callback_query"]["data"] == "role:driver" and vk["callback_query"]["from"]["id"] == "vk:7"
    mx = channels.from_max({"update_type": "message_created", "timestamp": 1,
                            "message": {"sender": {"user_id": 9}, "recipient": {"chat_id": 9}, "body": {"text": "/start", "mid": "m1"}}})
    assert mx["message"]["text"] == "/start" and mx["message"]["chat"]["id"] == "max:9"
    from app.telegram import kb

    m = {"text": "<b>Привет</b> &amp; добро", "kb": kb([("Водитель", "role:driver"), ("Кабинет", "app")])}
    body = channels.vk_payload(7, m)
    assert body["message"] == "Привет & добро" and '"type": "callback"' in body["keyboard"] and "open_link" in body["keyboard"]
    assert channels.max_payload(9, m)["attachments"][0]["payload"]["buttons"][0][0] == {"type": "callback", "text": "Водитель", "payload": "role:driver"}
    assert c.post("/api/v1/vk/callback", json={"type": "confirmation"}).status_code == 404  # not configured
    monkeypatch.setenv("VK_TOKEN", "t")
    monkeypatch.setenv("VK_SECRET", "s")
    monkeypatch.setenv("VK_CONFIRMATION", "abc123")
    assert c.post("/api/v1/vk/callback", json={"type": "confirmation"}).text == "abc123"
    event = {"type": "message_new", "event_id": "ev-test-1", "secret": "s", "object": {"message": {"from_id": 7, "peer_id": 7, "text": "/start"}}}
    assert c.post("/api/v1/vk/callback", json={**event, "secret": "wrong"}).status_code == 403
    assert c.post("/api/v1/vk/callback", json=event).text == "ok"
    assert c.post("/api/v1/vk/callback", json=event).text == "ok"  # VK retries: the second is not queued
    with db() as s:
        jobs = s.scalars(select(Job).where(Job.kind == "channel", Job.data["update"]["update_id"].astext == "vk-ev-test-1")).all()
        assert len(jobs) == 1


def test_personal_data_never_falls_back_to_the_cloud(monkeypatch):
    monkeypatch.setenv("PII_LOCAL", "1")
    monkeypatch.setenv("OLLAMA_URL", "http://127.0.0.1:9/v1")  # nothing listens: the local model is down
    with pytest.raises(RuntimeError, match="не отправляется в облако"):
        inbox.by_model("space", ["data:image/jpeg;base64,AA=="])
